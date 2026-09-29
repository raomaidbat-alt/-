"""Чтение CSV/XLSX и запись результата."""
import csv
import io
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import NEW_COLUMNS

TEXT_COLUMNS = ("инн", "кпп", "окпо", "огрн", "код филиала", "основной телефон")


def read_table(path):
    """(заголовки, строки). Первая непустая строка считается заголовком."""
    path = Path(path)
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        wb.close()
    else:
        rows = _read_csv(path.read_bytes())
    rows = [r for r in rows if any(v not in (None, "") for v in r)]
    if not rows:
        raise ValueError("файл пустой")
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    return headers, [[_cell(v) for v in r] for r in rows[1:]]


def _cell(v):
    # ИНН из Excel может прийти числом: 7707083893.0
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return v


def _read_csv(data):
    for enc in ("utf-8-sig", "cp1251"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        text = data.decode("utf-8", errors="replace")
    sample = text[:5000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=";,\t")
    except csv.Error:
        class dialect(csv.excel):
            delimiter = ";" if sample.count(";") >= sample.count(",") else ","
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def write_table(path, headers, rows):
    path = Path(path)
    if path.suffix.lower() == ".csv":
        with open(path, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.writer(f, delimiter=";")
            w.writerow(headers)
            w.writerows(rows)
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Обогащение"
    ws.append(headers)
    text_idx = [i for i, h in enumerate(headers) if str(h).lower().startswith(TEXT_COLUMNS)]
    for r in rows:
        row = ["" if v is None else v for v in r]
        for i in text_idx:
            if i < len(row) and row[i] != "":
                row[i] = str(row[i])
        ws.append(row)
    for i in text_idx:
        for c in ws[get_column_letter(i + 1)][1:]:
            c.number_format = "@"

    new_fill = PatternFill("solid", fgColor="FFF2CC")
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if h in NEW_COLUMNS:
            c.fill = new_fill
        width = max([len(str(h))] + [len(str(r[i - 1] or "")) for r in rows[:200] if i - 1 < len(r)])
        ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    wb.save(path)
